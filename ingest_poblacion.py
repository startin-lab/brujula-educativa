#!/usr/bin/env python3
"""
Brújula Educativa — Ingesta de población (DANE)
Fundación Startin

Responde la pregunta más básica de todas y la que el proyecto no tenía cómo
contestar: ¿cuánta gente vive aquí? Sin ese número, «45.000 menores en edad
escolar» no dice si el municipio es un pueblo donde los niños son la mitad o
una ciudad donde son una fracción.

QUÉ SE INGESTA (verificado el 20/09/2026 contra dane.gov.co)

  Proyecciones de población municipal 2018-2042, basadas en el Censo Nacional
  de Población y Vivienda 2018, actualizadas por el DANE el 30 de julio de 2025.
  Dos archivos Excel, porque el DANE no los publica por API:

    PPED-AreaMun-2018-2042_VP.xlsx          total, cabecera y rural por municipio y año
    PPED-AreaSexoEdadMun-2018-2042_VP.xlsx  lo mismo abierto por sexo y edad simple (0 a 100+)

  Del segundo solo se toman las edades escolares. Es un archivo de 130 MB con
  312 columnas: se lee en modo streaming y se descarta casi todo.

POR QUÉ DANE Y NO WIKIPEDIA

  Las cifras de población de Wikipedia son copias de estas mismas proyecciones,
  a veces de la versión anterior, sin decir de qué año son. Para una herramienta
  que promete fecha y fuente en cada cifra, la copia no sirve: se va al origen.

LO QUE HAY QUE SABER DE ESTOS ARCHIVOS

  1. La fila de encabezado es la 8 (las siete anteriores son título y notas).
     En el archivo por edad hay una fila 9 con el segundo nivel del encabezado
     («Total», «Hombres», «Mujeres», «Total 0 años», ...), y los datos empiezan
     en la 10.
  2. Cada municipio-año trae TRES filas: «Cabecera Municipal», «Centros Poblados
     y Rural Disperso» y «Total». Sumar sin filtrar duplica la población.
  3. El código municipal (MPIO) ya viene con cinco dígitos y como texto. Se
     rellena igual, por si un día deja de venir así.
  4. Son PROYECCIONES: el DANE las revisa. La cifra de 2026 hoy no será la
     cifra de 2026 dentro de dos años. Por eso viajan con el año y con la
     fecha de actualización del DANE.

Uso:
    python ingest_poblacion.py --salida ./data
    python ingest_poblacion.py --salida ./data --archivo-area ruta.xlsx --archivo-edad ruta.xlsx
"""

from __future__ import annotations

import argparse
import logging
import re
import sys
import tempfile
import time
from datetime import date
from pathlib import Path

import pandas as pd
import requests

LOG = logging.getLogger("poblacion")

BASE = "https://www.dane.gov.co/files/censo2018/proyecciones-de-poblacion/Municipal"
URL_AREA = f"{BASE}/PPED-AreaMun-2018-2042_VP.xlsx"
URL_EDAD = f"{BASE}/PPED-AreaSexoEdadMun-2018-2042_VP.xlsx"
ACTUALIZACION_DANE = "2025-07-30"     # «Actualizado el 30 de Julio de 2025», hoja PPED

TIMEOUT = 600
REINTENTOS = 3
ESPERA = 10

FILA_ENCABEZADO = 8
EDAD_MIN, EDAD_MAX_16, EDAD_MAX_18 = 5, 16, 18
# Se conservan los años desde el censo hasta un poco más allá de hoy: sirven
# para decir «la población escolar de este municipio cae desde 2019», que es
# una frase que cambia decisiones.
ANIO_DESDE = 2018
ANIO_HASTA = date.today().year + 5


def descargar(url: str, destino: Path) -> Path:
    ultimo: Exception | None = None
    for intento in range(1, REINTENTOS + 1):
        try:
            LOG.info("Descargando %s ...", url.rsplit("/", 1)[-1])
            with requests.get(url, stream=True, timeout=TIMEOUT) as r:
                r.raise_for_status()
                with open(destino, "wb") as f:
                    for trozo in r.iter_content(1 << 20):
                        f.write(trozo)
            tam = destino.stat().st_size
            if tam < 100_000:
                raise RuntimeError(f"archivo sospechosamente pequeño ({tam} bytes)")
            LOG.info("  %s MB", round(tam / 1e6, 1))
            return destino
        except Exception as exc:  # noqa: BLE001
            ultimo = exc
            LOG.warning("intento %s/%s falló: %s", intento, REINTENTOS, str(exc)[:160])
            if intento < REINTENTOS:
                time.sleep(ESPERA * intento)
    raise RuntimeError(f"No se pudo descargar {url}") from ultimo


def hoja_de_datos(wb) -> str:
    """La hoja grande. El nombre cambia entre archivos («PobMunicipalxÁrea»,
    «PobMunicipalxÁreaSexoEdad»); la de datos es la única con miles de filas."""
    candidatas = sorted(wb.worksheets, key=lambda ws: ws.max_row or 0, reverse=True)
    return candidatas[0].title


def codigo5(valor) -> str:
    return re.sub(r"\.0$", "", str(valor or "").strip()).zfill(5)


# --------------------------------------------------------------------------- #

def leer_area(ruta: Path) -> pd.DataFrame:
    """Total, cabecera y rural por municipio y año."""
    import openpyxl

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[hoja_de_datos(wb)]
    filas = ws.iter_rows(min_row=FILA_ENCABEZADO, values_only=True)
    encabezado = [str(c or "").strip().upper() for c in next(filas)]
    idx = {n: i for i, n in enumerate(encabezado)}
    for necesaria in ("DP", "DPNOM", "MPIO", "DPMP", "AÑO", "ÁREA GEOGRÁFICA", "TOTAL"):
        if necesaria not in idx:
            raise RuntimeError(f"El archivo por área no trae la columna {necesaria}: {encabezado}")

    registros: dict[tuple[str, int], dict] = {}
    for fila in filas:
        cod = fila[idx["MPIO"]]
        anio = fila[idx["AÑO"]]
        if cod is None or anio is None:
            continue
        try:
            anio = int(anio)
        except (TypeError, ValueError):
            continue
        if not (ANIO_DESDE <= anio <= ANIO_HASTA):
            continue
        area = str(fila[idx["ÁREA GEOGRÁFICA"]] or "").strip().lower()
        total = fila[idx["TOTAL"]]
        clave = (codigo5(cod), anio)
        reg = registros.setdefault(clave, {
            "cod_municipio": clave[0], "anio": anio,
            "cod_departamento": str(fila[idx["DP"]] or "").strip().zfill(2),
            "departamento_dane": str(fila[idx["DPNOM"]] or "").strip(),
            "municipio_dane": str(fila[idx["DPMP"]] or "").strip(),
        })
        if area.startswith("total"):
            reg["poblacion_total"] = total
        elif area.startswith("cabecera"):
            reg["poblacion_cabecera"] = total
        else:
            reg["poblacion_rural"] = total
    wb.close()

    df = pd.DataFrame(list(registros.values()))
    for c in ("poblacion_total", "poblacion_cabecera", "poblacion_rural"):
        if c not in df.columns:
            df[c] = pd.NA
        df[c] = pd.to_numeric(df[c], errors="coerce").astype("Int64")
    LOG.info("Área: %s municipios, años %s-%s", df["cod_municipio"].nunique(),
             int(df["anio"].min()), int(df["anio"].max()))
    return df


def leer_edades(ruta: Path) -> pd.DataFrame:
    """Población de 5 a 16 y de 5 a 18 años por municipio y año (ambos sexos)."""
    import openpyxl

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    ws = wb[hoja_de_datos(wb)]
    filas = ws.iter_rows(min_row=FILA_ENCABEZADO, values_only=True)
    primero = [str(c or "").strip().upper() for c in next(filas)]
    segundo = [str(c or "").strip() for c in next(filas)]
    idx = {n: i for i, n in enumerate(primero) if n}
    for necesaria in ("MPIO", "AÑO", "ÁREA GEOGRÁFICA"):
        if necesaria not in idx:
            raise RuntimeError(f"El archivo por edad no trae la columna {necesaria}")

    # Las columnas «Total N años» (ambos sexos, edad simple). El de 100 dice
    # «Total 100 años  y más»: no entra en ninguna franja escolar.
    col_edad: dict[int, int] = {}
    for i, nombre in enumerate(segundo):
        m = re.fullmatch(r"Total\s+(\d{1,3})\s+años.*", nombre)
        if m:
            col_edad[int(m.group(1))] = i
    faltan = [e for e in range(EDAD_MIN, EDAD_MAX_18 + 1) if e not in col_edad]
    if faltan:
        raise RuntimeError(f"El archivo por edad no trae las edades {faltan}")
    cols_16 = [col_edad[e] for e in range(EDAD_MIN, EDAD_MAX_16 + 1)]
    cols_18 = [col_edad[e] for e in range(EDAD_MIN, EDAD_MAX_18 + 1)]

    registros = []
    for fila in filas:
        cod, anio = fila[idx["MPIO"]], fila[idx["AÑO"]]
        if cod is None or anio is None:
            continue
        try:
            anio = int(anio)
        except (TypeError, ValueError):
            continue
        if not (ANIO_DESDE <= anio <= ANIO_HASTA):
            continue
        if not str(fila[idx["ÁREA GEOGRÁFICA"]] or "").strip().lower().startswith("total"):
            continue
        suma = lambda cols: sum(int(fila[c] or 0) for c in cols)  # noqa: E731
        registros.append({"cod_municipio": codigo5(cod), "anio": anio,
                          "poblacion_5_16_dane": suma(cols_16),
                          "poblacion_5_18_dane": suma(cols_18)})
    wb.close()
    df = pd.DataFrame(registros)
    LOG.info("Edades: %s municipios, años %s-%s", df["cod_municipio"].nunique(),
             int(df["anio"].min()), int(df["anio"].max()))
    return df


def unir(area: pd.DataFrame, edades: pd.DataFrame | None) -> pd.DataFrame:
    df = area
    if edades is not None and not edades.empty:
        df = df.merge(edades, on=["cod_municipio", "anio"], how="left")
        for c in ("poblacion_5_16_dane", "poblacion_5_18_dane"):
            df[c] = df[c].astype("Int64")
    df["fuente"] = "DANE — Proyecciones de población 2018-2042 (CNPV 2018)"
    df["actualizacion_dane"] = ACTUALIZACION_DANE
    return df.sort_values(["cod_municipio", "anio"]).reset_index(drop=True)


def validar(df: pd.DataFrame) -> None:
    """Lo mínimo para no publicar basura con cara de dato."""
    hoy = date.today().year
    del_anio = df[df["anio"] == hoy]
    if len(del_anio) < 1000:
        raise RuntimeError(f"Solo {len(del_anio)} municipios con población para {hoy}")
    total_pais = int(del_anio["poblacion_total"].sum())
    # Colombia ronda los 52-54 millones en esta década. Fuera de ese orden de
    # magnitud, lo que se leyó no es la columna que creemos.
    if not (45_000_000 < total_pais < 65_000_000):
        raise RuntimeError(f"La suma nacional de {hoy} da {total_pais:,}: columna equivocada")
    bog = del_anio[del_anio["cod_municipio"] == "11001"]
    if bog.empty or not (6_000_000 < int(bog.iloc[0]["poblacion_total"]) < 10_000_000):
        raise RuntimeError("Bogotá no aparece o su población no es verosímil")
    LOG.info("Validación: %s municipios, %s habitantes en %s", len(del_anio),
             f"{total_pais:,}".replace(",", "."), hoy)


def main() -> int:
    parser = argparse.ArgumentParser(description="Ingesta de población DANE")
    parser.add_argument("--salida", type=Path, default=Path("./data"))
    parser.add_argument("--archivo-area", type=Path, help="Excel por área ya descargado")
    parser.add_argument("--archivo-edad", type=Path, help="Excel por sexo y edad ya descargado")
    parser.add_argument("--sin-edades", action="store_true", help="omitir el archivo grande")
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args()

    logging.basicConfig(level=logging.DEBUG if args.verbose else logging.INFO,
                        format="%(levelname)-8s %(message)s")
    args.salida.mkdir(parents=True, exist_ok=True)

    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        area_xlsx = args.archivo_area or descargar(URL_AREA, tmp / "area.xlsx")
        area = leer_area(area_xlsx)

        edades = None
        if not args.sin_edades:
            try:
                edad_xlsx = args.archivo_edad or descargar(URL_EDAD, tmp / "edad.xlsx")
                edades = leer_edades(edad_xlsx)
            except Exception as exc:  # noqa: BLE001
                # El archivo grande es el que más falla (130 MB desde un
                # servidor lento). Sin él sigue habiendo habitantes; solo se
                # pierde la franja escolar DANE, que el MEN cubre en parte.
                LOG.error("Edades: %s — se publica sin franjas de edad", exc)

    df = unir(area, edades)
    validar(df)
    df.to_parquet(args.salida / "poblacion_municipios.parquet", index=False)

    hoy = date.today().year
    print("\n" + "=" * 52)
    print(f"Municipios              : {df['cod_municipio'].nunique():,}".replace(",", "."))
    print(f"Años                    : {int(df['anio'].min())}-{int(df['anio'].max())}")
    print(f"Habitantes {hoy} (país) : {int(df[df['anio'] == hoy]['poblacion_total'].sum()):,}".replace(",", "."))
    print(f"Franjas de edad         : {'sí' if edades is not None else 'NO'}")
    print("=" * 52)
    return 0


if __name__ == "__main__":
    sys.exit(main())
